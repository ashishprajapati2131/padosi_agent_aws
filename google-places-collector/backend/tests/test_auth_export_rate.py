import csv
import io
from datetime import datetime, timedelta, timezone

from app.services.errors import QUOTA_MESSAGE, GoogleServiceError
from app.services.export_service import render_csv, render_xlsx
from openpyxl import load_workbook


def test_register_login_me_and_refresh(api):
    registered = api.client.post(
        "/api/auth/register",
        json={"email": "Person@Example.com", "password": "password123"},
    )
    assert registered.status_code == 200, registered.text
    body = registered.json()
    assert body["user"]["email"] == "person@example.com"
    assert body["user"]["role"] == "user"
    assert "password" not in body["user"]
    me = api.client.get("/api/auth/me", headers={"Authorization": f"Bearer {body['access_token']}"})
    assert me.status_code == 200
    assert me.json()["email"] == "person@example.com"
    refreshed = api.client.post("/api/auth/refresh", json={"refresh_token": body["refresh_token"]})
    assert refreshed.status_code == 200
    assert refreshed.json()["access_token"]
    bad = api.client.post("/api/auth/login", json={"email": "person@example.com", "password": "wrong-password"})
    assert bad.status_code == 401
    assert bad.json()["detail"] == "Invalid email or password."
    duplicate = api.client.post(
        "/api/auth/register",
        json={"email": "person@example.com", "password": "password123"},
    )
    assert duplicate.status_code == 409


def test_user_cannot_open_admin(auth):
    denied = auth.client.get("/api/admin/stats", headers=auth.headers)
    assert denied.status_code == 403
    assert denied.json()["detail"] == "Admin access is required."
    cleanup = auth.client.post("/api/admin/retention/cleanup", headers=auth.headers)
    assert cleanup.status_code == 403


def test_admin_stats_and_retention(auth):
    from app.database.connection import get_session_factory
    from app.models import Place, SavedPlace, Search, User

    created = auth.client.post(
        "/api/search",
        json={"category": "Restaurant", "location": "Satellite, Ahmedabad"},
        headers=auth.headers,
    )
    assert created.status_code == 200, created.text
    keep_id = created.json()["search"]["id"]
    place_id = created.json()["results"][0]["id"]
    saved = auth.client.post("/api/saved-places", json={"place_id": place_id}, headers=auth.headers)
    assert saved.status_code == 200, saved.text

    auth.provider.calls.clear()
    auth.provider.pages = [[{"id": "other-place", "displayName": {"text": "Old Cafe"}, "formattedAddress": "Surat"}]]
    old = auth.client.post(
        "/api/search",
        json={"category": "Cafe", "location": "Surat"},
        headers=auth.headers,
    )
    assert old.status_code == 200, old.text
    old_id = old.json()["search"]["id"]

    db = get_session_factory()()
    user = db.query(User).filter(User.email == auth.email).one()
    user.role = "admin"
    aged = db.get(Search, old_id)
    aged.created_at = datetime.now(timezone.utc) - timedelta(days=120)
    db.commit()
    db.close()

    stats = auth.client.get("/api/admin/stats", headers=auth.headers)
    assert stats.status_code == 200, stats.text
    payload = stats.json()
    assert payload["total_users"] >= 1
    assert payload["total_searches"] >= 2
    assert payload["estimated_google_requests"] >= 1
    assert "dollar" not in str(payload).lower()

    usage = auth.client.get("/api/admin/api-usage", headers=auth.headers)
    assert usage.status_code == 200
    assert len(usage.json()["daily_searches"]) == 30

    cleaned = auth.client.post("/api/admin/retention/cleanup", headers=auth.headers)
    assert cleaned.status_code == 200, cleaned.text
    assert cleaned.json()["deleted_searches"] >= 1
    db = get_session_factory()()
    assert db.get(Search, old_id) is None
    assert db.get(Search, keep_id) is not None
    assert db.get(Place, place_id) is not None
    assert db.query(SavedPlace).count() == 1
    db.close()


def test_export_allowlist_omits_disabled_fields(monkeypatch):
    monkeypatch.setenv("EXPORT_DISABLED_FIELDS", "latitude,longitude")
    from app.config import reset_settings_cache

    reset_settings_cache()
    rows = [
        {
            "name": "River Hotel",
            "address": "Ahmedabad",
            "phone": "079",
            "website": "https://example.com",
            "rating": 4.5,
            "review_count": 10,
            "place_id": "abc",
            "latitude": 23.123456,
            "longitude": 72.5,
            "business_status": "OPERATIONAL",
            "maps_uri": "https://maps.google.com/?cid=1",
        }
    ]
    csv_text = render_csv(rows).decode("utf-8-sig")
    header = next(csv.reader(io.StringIO(csv_text)))
    assert "latitude" not in header
    assert "longitude" not in header
    assert "name" in header
    assert "23.123456" not in csv_text

    workbook = load_workbook(io.BytesIO(render_xlsx(rows)))
    sheet = workbook.active
    xlsx_header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert "latitude" not in xlsx_header
    assert "name" in xlsx_header
    reset_settings_cache()


def test_export_endpoint_uses_allowlist(auth, monkeypatch):
    created = auth.client.post(
        "/api/search",
        json={"category": "Office", "location": "SG Highway, Ahmedabad"},
        headers=auth.headers,
    )
    assert created.status_code == 200, created.text
    search_id = created.json()["search"]["id"]
    monkeypatch.setenv("EXPORT_DISABLED_FIELDS", "phone")
    from app.config import reset_settings_cache

    reset_settings_cache()
    exported = auth.client.post(
        "/api/exports",
        json={"search_id": search_id, "format": "csv"},
        headers=auth.headers,
    )
    assert exported.status_code == 200, exported.text
    assert "text/csv" in exported.headers["content-type"]
    text = exported.content.decode("utf-8-sig")
    header = next(csv.reader(io.StringIO(text)))
    assert "phone" not in header
    assert "name" in header
    listed = auth.client.get("/api/exports", headers=auth.headers)
    assert listed.status_code == 200
    assert listed.json()["total"] == 1


def test_rate_limit(auth, monkeypatch):
    monkeypatch.setenv("MAX_SEARCHES_PER_MINUTE", "2")
    monkeypatch.setenv("MAX_SEARCHES_PER_MINUTE_IP", "100")
    from app.config import reset_settings_cache
    from app.services.rate_limit import limiter

    reset_settings_cache()
    limiter.reset()
    assert _post(auth).status_code == 200
    assert _post(auth).status_code == 200
    blocked = _post(auth)
    assert blocked.status_code == 429
    assert "2 searches per minute" in blocked.json()["detail"]
    reset_settings_cache()


def test_bulk_preview_confirm_and_stop(auth):
    preview = auth.client.post(
        "/api/bulk-search",
        json={
            "csv_text": "category,location,radius\nHotel,Ahmedabad,5\nSchool,Ahmedabad,3\nHospital,Surat,4\n",
            "confirm": False,
        },
        headers=auth.headers,
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["confirm_required"] is True
    assert body["search_count"] == 3
    assert body["message"] == "You are about to perform 3 searches."
    assert auth.provider.calls == []

    auth.provider.error = None

    class StopAfterOne:
        def __init__(self):
            self.calls = []

        def search(self, request):
            self.calls.append(request)
            if len(self.calls) >= 2:
                raise GoogleServiceError("quota_exceeded", QUOTA_MESSAGE, 429, 429)
            from app.services.normalized import ProviderSearchPage

            return ProviderSearchPage(
                places=[],
                next_page_token=None,
                status_code=200,
                duration_ms=1,
                operation="searchText",
                result_count=0,
            )

        def get_details(self, place_id, field_mode):
            raise AssertionError("details")

    auth.provider = StopAfterOne()
    confirmed = auth.client.post(
        "/api/bulk-search",
        json={
            "rows": [
                {"category": "Hotel", "location": "Ahmedabad", "radius_km": 5},
                {"category": "School", "location": "Ahmedabad", "radius_km": 3},
                {"category": "Hospital", "location": "Surat", "radius_km": 4},
            ],
            "confirm": True,
        },
        headers=auth.headers,
    )
    assert confirmed.status_code == 200, confirmed.text
    result = confirmed.json()
    assert result["stopped"] is True
    assert result["failed_row"] == 2
    assert result["stop_reason"] == QUOTA_MESSAGE
    assert len(result["results"]) == 2
    assert len(auth.provider.calls) == 2


def _post(auth):
    return auth.client.post(
        "/api/search",
        json={"category": "Gym", "location": "Navrangpura"},
        headers=auth.headers,
    )
